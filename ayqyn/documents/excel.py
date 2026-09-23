"""Bounded XLSX extraction with cell provenance; never evaluates formulas."""
import hashlib
import io
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

MAX_UPLOAD_BYTES = 10_000_000
MAX_UNCOMPRESSED_BYTES = 32_000_000
MAX_SHEETS = 32
MAX_GRID_CELLS = 200_000
MAX_NONEMPTY_CELLS = 10_000
MAX_TEXT_CHARS = 250_000
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def _bounds(reference):
    bounds = range_boundaries(reference)
    left, top, right, bottom = bounds
    if (any(v is None for v in bounds) or min(bounds) < 1 or right > 16384
            or bottom > 1048576 or right < left or bottom < top
            or (right-left+1)*(bottom-top+1) > MAX_GRID_CELLS):
        raise ValueError('XLSX range exceeds supported bounds; nothing was truncated.')
    return bounds


def _preflight(data):
    """Bound XML and merged-cell expansion before openpyxl loads the workbook."""
    with ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        names = [item.filename for item in items]
        if len(items) > 2000 or len(names) != len(set(names)):
            raise ValueError('Unsupported XLSX archive structure.')
        if sum(item.file_size for item in items) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError('XLSX expanded size exceeds 32 MB.')
        if any(item.flag_bits & 1 for item in items) or any('vbaProject' in n for n in names):
            raise ValueError('Encrypted or macro-enabled workbooks are not supported.')
        sheets = [n for n in names if n.startswith('xl/worksheets/') and n.endswith('.xml')]
        if not sheets or len(sheets) > MAX_SHEETS:
            raise ValueError('XLSX must contain 1..32 worksheets.')
        for name in names:
            if not name.endswith(('.xml', '.rels')):
                continue
            raw = archive.read(name)
            if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                raise ValueError('XML entities are not supported.')
            if name in sheets:
                root = ET.fromstring(raw)
                cells = root.findall('.//'+NS+'c')
                if len(cells) > MAX_GRID_CELLS:
                    raise ValueError('XLSX worksheet has too many cells.')
                for cell in cells:
                    _bounds(cell.attrib['r'])
                expanded = 0
                for merge in root.findall('.//'+NS+'mergeCell'):
                    left, top, right, bottom = _bounds(merge.attrib['ref'])
                    expanded += (right-left+1)*(bottom-top+1)
                if expanded > MAX_GRID_CELLS:
                    raise ValueError('XLSX merged ranges exceed supported bounds.')
        return [n for n in names if n.startswith(('xl/drawings/', 'xl/charts/', 'xl/embeddings/'))
                and not n.endswith('/')]


def _text(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    return str(value)


def parse_xlsx(data: bytes, filename: str) -> dict:
    if Path(filename).suffix.lower() != '.xlsx':
        raise ValueError('Supported Excel format: .xlsx; .xls/.xlsm require conversion.')
    if not isinstance(data, bytes) or not data or len(data) > MAX_UPLOAD_BYTES:
        raise ValueError('Provide a non-empty XLSX up to 10 MB.')
    workbook = cached = None
    try:
        unread_objects = _preflight(data)
        workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ValueError('XLSX exceeds 32 worksheets.')
        paragraphs, sheets, unread_cells = [], [], []
        warnings = ['XLSX: text is extracted from cells, not rendered pages. Numeric/date values '
                    'are normalized; number_format is retained. Formulas are not recalculated. '
                    'Row/merge/header context does not establish ownership.']
        text_chars = grid_cells = formula_count = 0
        for sheet_index, sheet in enumerate(workbook.worksheets, 1):
            prefix = f's{sheet_index}'
            grid_cells += sheet.max_row*sheet.max_column
            if grid_cells > MAX_GRID_CELLS:
                raise ValueError('XLSX grid exceeds 200,000 cells; nothing was truncated.')
            merges = [(str(r), _bounds(str(r))) for r in sheet.merged_cells.ranges]
            tables = [(table.name, _bounds(table.ref), table.headerRowCount or 0)
                      for table in sheet.tables.values()]
            sheet_info = {'name':sheet.title, 'index':sheet_index, 'state':sheet.sheet_state,
                          'merged_ranges':[ref for ref,_ in merges],
                          'tables':[{'name':t.name,'range':t.ref,'header_rows':t.headerRowCount or 0}
                                    for t in sheet.tables.values()]}
            sheets.append(sheet_info)
            sheet_fragments = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    formula = None
                    value = cell.value
                    cached_value = None
                    if cell.data_type == 'f':
                        formula_count += 1
                        formula = value if isinstance(value,str) else getattr(value,'text',None)
                        cached_cell = cached[sheet.title][cell.coordinate]
                        cached_value = cached_cell.value
                        if cached_value is None or cached_cell.data_type == 'e':
                            unread_cells.append(f"'{sheet.title}'!{cell.coordinate}")
                        # A missing cache is represented as formula text, never a fabricated value.
                        value = cached_value if cached_value is not None else formula or '[unsupported formula]'
                    if cell.data_type == 'e':
                        unread_cells.append(f"'{sheet.title}'!{cell.coordinate}")
                    text = _text(value)
                    if not text.strip():
                        continue
                    text_chars += len(text)
                    if text_chars > MAX_TEXT_CHARS or len(paragraphs) >= MAX_NONEMPTY_CELLS:
                        raise ValueError('XLSX text/cell limit exceeded; nothing was truncated.')
                    table_name, origin_row, is_header = 'grid', 1, False
                    for name,(left,top,right,bottom),header_rows in tables:
                        if left <= cell.column <= right and top <= cell.row <= bottom:
                            table_name, origin_row = name, top
                            is_header = cell.row < top+header_rows
                            break
                    table_id = f'{prefix}:table:{table_name}'
                    own_merge = next((ref for ref,(left,top,right,bottom) in merges
                                      if cell.row==top and cell.column==left),None)
                    # Keep geometric context separate from semantic parent/owner assignment.
                    context_ids = [f'{prefix}!{sheet.cell(top,left).coordinate}'
                                   for _,(left,top,right,bottom) in merges
                                   if top <= cell.row <= bottom and
                                   (cell.row,cell.column)!=(top,left)]
                    fragment = {'id':f'{prefix}!{cell.coordinate}', 'text':text,
                                'section':f"'{sheet.title}'!{cell.coordinate}",
                                'sheet':sheet.title,'sheet_index':sheet_index,'cell':cell.coordinate,
                                'cell_range':own_merge or cell.coordinate,'page':None,'printed_page':None,
                                'block_type':'table_cell','parent_id':None,
                                'table_id':table_id,'row_id':f'{table_id}:r{cell.row}',
                                'row_index':cell.row-origin_row,'column_index':cell.column-1,
                                'row_number':cell.row,'is_header':is_header,
                                'context_fragment_ids':context_ids,
                                'sheet_state':sheet.sheet_state,
                                'hidden_row':bool(sheet.row_dimensions[cell.row].hidden),
                                'number_format':cell.number_format,'value_type':cell.data_type,
                                'formula':formula,'formula_cached':cell.data_type=='f' and cached_value is not None}
                    paragraphs.append(fragment)
                    sheet_fragments.append(fragment)
            known={p['id'] for p in sheet_fragments}
            for p in sheet_fragments:
                p['context_fragment_ids']=[key for key in p['context_fragment_ids'] if key in known]
            if sheet.sheet_state != 'visible':
                warnings.append(f'Hidden sheet included: {sheet.title}.')
        if not paragraphs:
            raise ValueError('XLSX has no non-empty extractable cells.')
        if formula_count:
            warnings.append('Formula caches may be stale; no formulas or external links were executed.')
        if unread_cells:
            warnings.append(f'{len(unread_cells)} cells have missing formula results or Excel errors; extraction is partial.')
        if unread_objects:
            warnings.append('Drawings/charts/embedded objects in XLSX were not interpreted; no XLSX OCR/VLM was performed.')
        return {'name':filename,'sha256':hashlib.sha256(data).hexdigest(),'format':'xlsx',
                'paragraphs':paragraphs,'sheets':sheets,'sheet_count':len(sheets),'page_count':None,
                'warnings':warnings,'unread_cells':unread_cells,'unread_objects':unread_objects,
                'extraction_incomplete':bool(unread_cells or unread_objects)}
    except (BadZipFile, ET.ParseError, KeyError, TypeError, AttributeError, OSError, IndexError) as exc:
        raise ValueError('Malformed or unsupported XLSX workbook.') from exc
    finally:
        if workbook is not None:
            workbook.close()
        if cached is not None:
            cached.close()
